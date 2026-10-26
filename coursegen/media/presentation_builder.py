"""Build PowerPoint presentations and their visual assets."""

from __future__ import annotations

import io
import logging
from pathlib import Path
from typing import Dict, Optional
from urllib.parse import urlparse

import requests
from bing_image_urls import bing_image_urls
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from coursegen.config import Settings, get_settings
from coursegen.errors import ValidationError
from coursegen.media.fonts import load_font
from coursegen.models import LectureMap, PresentationStyle
from coursegen.utils import safe_filename

logger = logging.getLogger(__name__)

_IMAGE_EXTENSIONS = (".jpg", ".jpeg", ".png", ".gif", ".bmp")
_TEMPLATE_FILES: Dict[str, Optional[str]] = {
    "Default": None,
    "Dark": "dark.png",
    "Fancy": "Fancy.png",
    "History": "history.png",
    "Modern": "modern.jpg",
    "White": "white.png",
}


class TemplateLibrary:
    def __init__(self, settings: Settings | None = None) -> None:
        self._settings = settings or get_settings()

    @property
    def names(self) -> list[str]:
        return list(_TEMPLATE_FILES)

    def path_for(self, template: str) -> Optional[Path]:
        filename = _TEMPLATE_FILES.get(template)
        if not filename:
            return None
        path = self._settings.backgrounds_dir / filename
        return path if path.is_file() else None


class FontPreviewer:
    @staticmethod
    def render(font_name: str, font_size: int) -> Image.Image:
        image = Image.new("RGB", (300, 116), color="white")
        draw = ImageDraw.Draw(image)
        font = load_font(font_name, font_size)
        text = f"{font_name} 1 2 3"
        left, top, right, bottom = draw.textbbox((0, 0), text, font=font)
        draw.text(
            ((300 - (right - left)) / 2, (100 - (bottom - top)) / 2),
            text,
            fill="black",
            font=font,
        )
        return image


class ImageFetcher:
    """Download one illustration per slide topic."""

    def __init__(self, settings: Settings | None = None) -> None:
        self._settings = settings or get_settings()
        self._settings.images_dir.mkdir(parents=True, exist_ok=True)

    def save(self, image_content: bytes, title: str) -> Optional[Path]:
        try:
            with Image.open(io.BytesIO(image_content)) as image:
                image_format = (image.format or "").upper()
        except Exception:
            logger.warning("Downloaded content for %s is not a valid image.", title)
            return None

        extension = {
            "JPEG": "jpg",
            "PNG": "png",
            "GIF": "gif",
            "BMP": "bmp",
        }.get(image_format)
        if extension is None:
            logger.info("Skipping unsupported image format %s for %s.", image_format, title)
            return None

        path = self._settings.images_dir / f"{safe_filename(title, 'slide')}.{extension}"
        path.write_bytes(image_content)
        return path

    def scrape(self, slides: LectureMap) -> Dict[str, Dict[str, str]]:
        final_images: Dict[str, Dict[str, str]] = {}
        for lecture, topics in slides.items():
            images: Dict[str, str] = {}
            for topic in topics:
                try:
                    urls = bing_image_urls(topic, limit=3)
                except Exception as exc:
                    logger.warning("Image search failed for %s: %s", topic, exc)
                    urls = []
                for url in urls:
                    if urlparse(url).path.lower().endswith(_IMAGE_EXTENSIONS):
                        images[topic] = url
                        break
            final_images[lecture] = images
        return final_images


class PresentationBuilder:
    def __init__(
        self,
        settings: Settings | None = None,
        fetcher: ImageFetcher | None = None,
    ) -> None:
        self._settings = settings or get_settings()
        self._fetcher = fetcher or ImageFetcher(self._settings)

    def create(self, slides: LectureMap, style: PresentationStyle, name: str) -> Path:
        self._settings.presentations_dir.mkdir(parents=True, exist_ok=True)
        presentation = Presentation()
        image_urls = self._fetcher.scrape(slides)

        for lecture_title, topics in slides.items():
            self._add_slide(
                presentation,
                title=lecture_title,
                content="",
                image_url=None,
                style=style,
                is_lecture=True,
            )
            for topic_title, content_lines in topics.items():
                lines = [content_lines] if isinstance(content_lines, str) else list(content_lines)
                bullets = [
                    f"\u2022{line[1:]}" if line.startswith("-") else f"\u2022{line}"
                    for line in lines
                    if line.strip()
                ]
                for start in range(0, len(bullets), 3):
                    self._add_slide(
                        presentation,
                        title=topic_title.replace("?", "").replace("*", ""),
                        content="\n\n".join(bullets[start : start + 3]),
                        image_url=image_urls.get(lecture_title, {}).get(topic_title),
                        style=style,
                    )

        output_path = self._settings.presentations_dir / f"powerpoint_{safe_filename(name, 'course')}.pptx"
        presentation.save(str(output_path))
        logger.info("Created presentation %s", output_path.name)
        return output_path

    def _add_slide(
        self,
        presentation: Presentation,
        title: str,
        content: str,
        image_url: Optional[str],
        style: PresentationStyle,
        is_lecture: bool = False,
    ) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])

        if style.background_path:
            background = Path(style.background_path)
            if background.is_file():
                slide.shapes.add_picture(
                    str(background),
                    0,
                    0,
                    width=presentation.slide_width,
                    height=presentation.slide_height,
                )

        title_height = Inches(1.1)
        title_left = Inches(0.2)
        title_width = presentation.slide_width - Inches(2)
        title_top = Inches(0.5)
        if is_lecture:
            title_left = Inches(3)
            title_top = (presentation.slide_height - title_height) / 2
            title_font_size = 60
        else:
            title_font_size = style.title_font_size

        self._add_text_box(
            slide,
            title,
            left=title_left,
            top=title_top,
            width=title_width,
            height=title_height,
            font_size=title_font_size,
            color=style.title_color,
            font_name=style.font_name,
        )

        if not content:
            return

        has_image = False
        if image_url:
            image_path = self._download_image(image_url, title)
            if image_path is not None:
                slide.shapes.add_picture(
                    str(image_path),
                    left=Inches(6),
                    top=Inches(2.25),
                    width=Inches(4),
                    height=Inches(4),
                )
                has_image = True

        width = presentation.slide_width - Inches(5 if has_image else 1)
        self._add_text_box(
            slide,
            content,
            left=Inches(0.6),
            top=Inches(1.5),
            width=width,
            height=presentation.slide_height - Inches(2.5),
            font_size=style.content_font_size,
            color=style.content_color,
            font_name=style.font_name,
        )

    def _download_image(self, url: str, title: str) -> Optional[Path]:
        try:
            response = requests.get(url, timeout=20)
            response.raise_for_status()
        except requests.RequestException as exc:
            logger.warning("Could not download image for %s: %s", title, exc)
            return None
        try:
            return self._fetcher.save(response.content, title)
        except OSError as exc:
            logger.warning("Could not save image for %s: %s", title, exc)
            return None

    @staticmethod
    def _add_text_box(
        slide: object,
        text: str,
        left: int,
        top: int,
        width: int,
        height: int,
        font_size: int,
        color: str,
        font_name: str,
    ) -> None:
        box = slide.shapes.add_textbox(left, top, width, height)
        frame = box.text_frame
        frame.word_wrap = True
        paragraph = frame.paragraphs[0]
        paragraph.text = text
        paragraph.font.size = Pt(font_size)
        paragraph.font.color.rgb = _hex_to_rgb(color)
        paragraph.font.name = font_name


def _hex_to_rgb(value: str) -> RGBColor:
    cleaned = value.lstrip("#")
    if len(cleaned) != 6:
        raise ValidationError(f"Invalid colour value: {value}")
    return RGBColor(*(int(cleaned[index : index + 2], 16) for index in (0, 2, 4)))
